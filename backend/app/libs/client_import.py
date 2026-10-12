"""Import clients from a spreadsheet (Comment #11): parse, map, validate.

Pure functions only -- no database. ``app.apis.clients`` feeds in the file
and the names already taken, and creates the rows this module marks
"create" through the same INSERT as the New Client form.

The template's headers are the user's, exactly as given (2026-10-12) --
the Christmas client sheet's own columns:

    First | Last | ADDRESS | CITY | ST | ZIP | Area | Zone | PHONE | EMAIL

Area and Zone are the install schedule's ("North" / "The Woodlands") and
are saved on the client (clients.area / clients.zone, migrations/023).

Headers match ignoring case, spaces and punctuation. Still read when
present: the New Client form's long labels (First name, Last name, Street
address, Business name, Location, Install time preference, Additional
contact label/phone/email, General notes) and common Odoo export headers
(``HEADER_ALIASES``):

    Odoo header                         goes into
    Name / Display Name / Contact Name  First + Last name for a person, or the
                                        Business name when the row is a
                                        company ("Is a Company" / "Company
                                        Type") or reads like one
    Company Name / Related Company      Business name (the Name is then the
                                        business's contact person)
    Street                              ADDRESS
    Street2                             added to ADDRESS
    Zip / Zip Code / Postal Code        ZIP
    State ("Texas (US)")                ST, as TX
    Mobile                              PHONE, or an additional contact
                                        "Mobile" when PHONE is filled too
    Notes / Internal Notes / Comment    General notes
    Email, Phone, City                  EMAIL, PHONE, CITY

Odoo's plain "Company" column is NOT read: on an Odoo contact it is the
Odoo company that owns the record (ours), not the client's business. Any
other column (Country, Tags, External ID...) is listed as ignored.

Rows are never merged into existing clients. A row whose name (composed by
``client_names.compose_name``, the same rule as the form) is already a
client's name -- or one of its former names -- is a duplicate and skipped,
as is a second row with the same name in the file.
"""

from __future__ import annotations

import base64
import csv
import io
import re
from datetime import date, datetime
from typing import Iterable, Optional

from app.libs import client_names

# ─── the template ───────────────────────────────────────────────────────────

#: (field key, header): the user's columns, exactly as given (2026-10-12),
#: which are the Christmas client sheet's own headers.
TEMPLATE_FIELDS: tuple[tuple[str, str], ...] = (
    ("first_name", "First"),
    ("last_name", "Last"),
    ("street", "ADDRESS"),
    ("city", "CITY"),
    ("state", "ST"),
    ("zip", "ZIP"),
    ("area", "Area"),
    ("zone", "Zone"),
    ("phone", "PHONE"),
    ("email", "EMAIL"),
)
TEMPLATE_HEADERS: tuple[str, ...] = tuple(h for _, h in TEMPLATE_FIELDS)

#: Columns still read when present, though not in the template: the New
#: Client form's other boxes, under the form's labels.
OPTIONAL_FIELDS: tuple[tuple[str, str], ...] = (
    ("company", "Business name"),
    ("site", "Location"),
    ("time_preference", "Install time preference"),
    ("contact_label", "Additional contact label"),
    ("contact_phone", "Additional contact phone"),
    ("contact_email", "Additional contact email"),
    ("notes", "General notes"),
)
FIELD_LABEL = {**dict(TEMPLATE_FIELDS), **dict(OPTIONAL_FIELDS)}

#: Fake rows for "Download example": three people, with Area / Zone as the
#: schedule zones their ZIP (scheduler/prep.py ZIP_ZONE).
EXAMPLE_ROWS: tuple[dict, ...] = (
    {"first_name": "Jamie", "last_name": "Rivera", "street": "123 Maple Ln", "city": "The Woodlands",
     "state": "TX", "zip": "77380", "area": "North", "zone": "The Woodlands",
     "phone": "(281) 555-0142", "email": "jamie.rivera@example.com"},
    {"first_name": "Dana", "last_name": "Lee", "street": "1418 Columbia St", "city": "Houston",
     "state": "TX", "zip": "77008", "area": "Central", "zone": "Heights/Near Northside",
     "phone": "(713) 555-0110", "email": "dana.lee@example.com"},
    {"first_name": "Marcus", "last_name": "Hale", "street": "2207 Prairie Rose Dr", "city": "Katy",
     "state": "TX", "zip": "77450", "area": "West", "zone": "Katy/Fulshear",
     "phone": "(832) 555-0199", "email": "marcus.hale@example.com"},
)

#: Extra keys only an alias can fill (no template column of their own).
_EXTRA_KEYS = ("name", "mobile", "street2", "is_company", "company_type")

#: Accepted header spellings per key, beyond the template header itself.
HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "first_name": ("first name", "firstname", "given name"),
    "last_name": ("last name", "lastname", "surname", "family name"),
    "company": ("business", "company name", "related company", "parent name",
                "parent company", "organization", "organisation"),
    "site": ("site", "property", "location name"),
    "email": ("e-mail", "email address", "e-mail address"),
    "phone": ("phone number", "telephone", "tel", "work phone", "home phone"),
    "street": ("street address", "street", "address line 1", "street 1", "street1"),
    "city": ("town",),
    "state": ("state", "state/province", "province", "state name", "region"),
    "area": ("schedule area",),
    "zone": ("schedule zone",),
    "zip": ("zip", "zip code", "zipcode", "postal code", "postcode", "post code"),
    "time_preference": ("time preference", "install time", "preferred time"),
    "contact_label": ("additional contact", "additional contact name"),
    "contact_phone": (),
    "contact_email": (),
    "notes": ("notes", "note", "internal notes", "internal note", "comment", "comments"),
    "name": ("name", "display name", "full name", "contact name", "client name",
             "customer", "customer name", "client", "partner", "contact"),
    "mobile": ("mobile", "mobile phone", "cell", "cell phone", "mobile number"),
    "street2": ("street2", "street 2", "address line 2", "address 2", "suite"),
    "is_company": ("is a company", "is company"),
    "company_type": ("company type",),
}

#: For the dialog's "From Odoo?" list: what the common Odoo headers become.
ODOO_ALIAS_DOC: tuple[tuple[str, str], ...] = (
    ("Name", "First + Last (or Business name for a company)"),
    ("Company Name / Related Company", "Business name (the Name becomes its contact person)"),
    ("Is a Company / Company Type", "says whether Name is a business"),
    ("Street", "ADDRESS"),
    ("Street2", "added to ADDRESS"),
    ("Zip", "ZIP"),
    ("State (e.g. Texas (US))", "ST, as TX"),
    ("Mobile", "PHONE, or an additional contact when PHONE is filled"),
    ("Notes / Internal Notes", "General notes"),
    ("Email, Phone, City", "EMAIL, PHONE, CITY"),
)

MAX_ROWS = 2000
MAX_BYTES = 3 * 1024 * 1024

STATUS_CREATE = "create"
STATUS_DUPLICATE = "duplicate"
STATUS_ERROR = "error"


class ImportFileError(ValueError):
    """The file as a whole can't be read (the API turns it into a 400)."""


# ─── headers ────────────────────────────────────────────────────────────────


def _norm_header(h: Optional[str]) -> str:
    h = (h or "").replace("\ufeff", "").strip().lower()
    h = re.sub(r"[_\-]+", " ", h)
    h = re.sub(r"[^a-z0-9/ ]+", "", h)
    return re.sub(r"\s+", " ", h).strip()


def _build_alias_index() -> dict[str, str]:
    index: dict[str, str] = {}
    for key, header in TEMPLATE_FIELDS + OPTIONAL_FIELDS:
        index[_norm_header(header)] = key
    for key, aliases in HEADER_ALIASES.items():
        for a in aliases:
            index.setdefault(_norm_header(a), key)
    return index


_ALIAS_INDEX = _build_alias_index()


def map_headers(headers: Iterable[str]) -> tuple[dict[int, str], list[str]]:
    """{column index: field key} for the headers we know, plus the headers
    we ignore. The first column for a key wins; a repeat is ignored."""
    mapping: dict[int, str] = {}
    ignored: list[str] = []
    used: set[str] = set()
    for i, raw in enumerate(headers):
        key = _ALIAS_INDEX.get(_norm_header(raw))
        if key and key not in used:
            mapping[i] = key
            used.add(key)
        elif (raw or "").strip():
            ignored.append(raw.strip())
    return mapping, ignored


# ─── reading the file ───────────────────────────────────────────────────────


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def read_rows(filename: str, data: bytes) -> list[list[str]]:
    """Every row of the file as text cells, header row first."""
    if not data:
        raise ImportFileError("That file is empty.")
    if len(data) > MAX_BYTES:
        raise ImportFileError("That file is too big (3 MB at most). Split it into smaller files.")
    name = (filename or "").lower()
    if data[:2] == b"PK" or name.endswith((".xlsx", ".xlsm")):
        return _read_xlsx(data)
    if name.endswith(".xls") or data[:4] == b"\xd0\xcf\x11\xe0":
        raise ImportFileError("Old .xls files can't be read. Save it as CSV (or .xlsx) and try again.")
    text = _decode(data)
    first_line = text.split("\n", 1)[0]
    delimiter = max((",", ";", "\t"), key=first_line.count)
    return [list(r) for r in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any unreadable workbook is the same answer
        raise ImportFileError("Couldn't open that spreadsheet. Save it as CSV and try again.") from exc
    try:
        ws = wb.active
        return [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def decode_upload(content_base64: Optional[str]) -> bytes:
    try:
        return base64.b64decode(content_base64 or "", validate=False)
    except Exception as exc:  # noqa: BLE001
        raise ImportFileError("Couldn't read that file.") from exc


# ─── cleaning one row ───────────────────────────────────────────────────────

_EMAIL = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[A-Za-z]{2,}$")
_EXT = re.compile(r"\s*(?:x|ext\.?|extension|#)\s*\d+\s*$", re.I)

US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR", "california": "CA",
    "colorado": "CO", "connecticut": "CT", "delaware": "DE", "district of columbia": "DC",
    "florida": "FL", "georgia": "GA", "hawaii": "HI", "idaho": "ID", "illinois": "IL",
    "indiana": "IN", "iowa": "IA", "kansas": "KS", "kentucky": "KY", "louisiana": "LA",
    "maine": "ME", "maryland": "MD", "massachusetts": "MA", "michigan": "MI", "minnesota": "MN",
    "mississippi": "MS", "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK", "oregon": "OR",
    "pennsylvania": "PA", "rhode island": "RI", "south carolina": "SC", "south dakota": "SD",
    "tennessee": "TN", "texas": "TX", "utah": "UT", "vermont": "VT", "virginia": "VA",
    "washington": "WA", "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}

TIME_PREFERENCE_WORDS = {
    "": None, "no preference": None, "none": None, "any": None, "anytime": None, "flexible": None,
    "morning": "morning", "early morning": "morning", "early": "morning", "am": "morning",
    "afternoon": "afternoon", "pm": "afternoon", "midday": "afternoon",
    "late": "late", "evening": "late", "night": "late", "late afternoon": "late",
}

_TRUE = {"true", "yes", "y", "1", "x", "company", "is a company"}


def _clean(v: Optional[str]) -> str:
    return re.sub(r"\s+", " ", (v or "").replace("\u00a0", " ")).strip()


def clean_state(v: str) -> str:
    s = re.sub(r"\s*\([A-Za-z]{2,3}\)\s*$", "", _clean(v))  # Odoo: "Texas (US)"
    if len(s) == 2 and s.isalpha():
        return s.upper()
    return US_STATES.get(s.lower(), s)


def clean_zip(v: str) -> str:
    z = _clean(v)
    if re.fullmatch(r"\d+\.0+", z):  # a spreadsheet's number: "77380.0"
        z = z.split(".")[0]
    if re.fullmatch(r"\d{3,4}", z):  # a spreadsheet dropped the leading zero
        z = z.zfill(5)
    return z


def phone_problem(v: str) -> Optional[str]:
    if not v:
        return None
    digits = re.sub(r"\D", "", _EXT.sub("", v))
    if re.search(r"[A-Za-z]", _EXT.sub("", v)) or not 7 <= len(digits) <= 15:
        return f'phone "{v}" doesn\'t look like a phone number'
    return None


def email_problem(v: str) -> Optional[str]:
    if v and not _EMAIL.match(v):
        return f'email "{v}" doesn\'t look like an email address'
    return None


def _truthy(v: str) -> bool:
    return _clean(v).lower() in _TRUE


def _title_if_lower(v: Optional[str]) -> Optional[str]:
    """"jamie" -> "Jamie"; anything with a capital is left as typed."""
    if v and v == v.lower() and any(c.isalpha() for c in v):
        return " ".join(w[:1].upper() + w[1:] for w in v.split(" "))
    return v


def _looks_like_business(text: str) -> bool:
    words = set(client_names._words(text))
    return bool(words & (client_names._BUSINESS_WORDS - client_names._TITLES)) or bool(
        re.search(r"\d|&|\b(llc|inc|ltd)\b", text, re.I))


def _split_full_name(full: str, notes: list[str]) -> dict:
    """An Odoo "Name" for a row with no First/Last/Business columns filled."""
    if full.count(",") == 1 and not _looks_like_business(full):
        last, first = (_clean(x) for x in full.split(",", 1))
        if last and first:
            return {"last_name": last, "first_name": first}
    words = full.split(" ")
    if 2 <= len(words) <= 3 and "," not in full and not _looks_like_business(full):
        first, last = " ".join(words[:-1]), words[-1]
        notes.append(f'Split "{full}" into First "{first}", Last "{last}"')
        return {"first_name": first, "last_name": last}
    notes.append(f'Read "{full}" as a business name')
    return {"company": full}


def build_row(cells: dict) -> tuple[dict, list[str], list[str]]:
    """Map one row's {key: text} to the create fields.

    Returns (fields, problems, notes): problems make the row an error,
    notes are what the import guessed (shown, not blocking)."""
    c = {k: _clean(v) for k, v in cells.items()}
    problems: list[str] = []
    notes: list[str] = []

    first, last, company, site = (c.get(k, "") for k in ("first_name", "last_name", "company", "site"))
    full = c.get("name", "")
    is_company = _truthy(c.get("is_company", "")) or c.get("company_type", "").lower() == "company"
    if full and not (first or last):
        if is_company and not company:
            company = full
        elif company:
            # Odoo's display name for a business's contact: "Business, Person".
            person = full[len(company) + 1:].strip() if full.lower().startswith(company.lower() + ",") else full
            if person and person.lower() != company.lower():
                split = _split_full_name(person, [])  # the contact note says it
                if "company" in split:  # not a person's name: leave it out
                    notes.append(f'Left "{person}" out (not a person\'s name)')
                else:
                    first, last = split.get("first_name", ""), split.get("last_name", "")
        else:
            split = _split_full_name(full, notes)
            first, last, company = split.get("first_name", ""), split.get("last_name", ""), split.get("company", "")

    titled_first, titled_last = _title_if_lower(first), _title_if_lower(last)
    if (titled_first, titled_last) != (first, last):
        notes.append("Capitalized the first/last name")
        first, last = titled_first, titled_last

    parts = None
    try:
        parts = client_names.clean_parts(first, last, company, site)
    except client_names.NamePartsError as exc:
        problems.append(str(exc))

    street = c.get("street", "")
    if c.get("street2"):
        street = f"{street}, {c['street2']}" if street else c["street2"]

    phone, mobile = c.get("phone", ""), c.get("mobile", "")
    contacts: list[dict] = []
    if mobile and not phone:
        phone = mobile
    elif mobile and re.sub(r"\D", "", mobile) != re.sub(r"\D", "", phone):
        contacts.append({"label": "Mobile", "phone": mobile, "email": None})

    label, cphone, cemail = (c.get(k, "") for k in ("contact_label", "contact_phone", "contact_email"))
    if label or cphone or cemail:
        contacts.append({"label": label or "Contact", "phone": cphone or None, "email": cemail or None})

    for value, problem in ((c.get("email", ""), email_problem), (phone, phone_problem)):
        p = problem(value)
        if p:
            problems.append(p)
    for contact in contacts:
        for value, problem in ((contact["phone"] or "", phone_problem), (contact["email"] or "", email_problem)):
            p = problem(value)
            if p:
                problems.append(f"additional contact {p}")

    raw_pref = c.get("time_preference", "")
    pref_key = raw_pref.lower()
    if pref_key not in TIME_PREFERENCE_WORDS:
        problems.append(f'Install time preference "{raw_pref}" should be Morning, Afternoon, Late or blank')
        pref = None
    else:
        pref = TIME_PREFERENCE_WORDS[pref_key]

    fields = {
        "parts": parts,
        "name": client_names.compose_name(**parts) if parts else "",
        "email": c.get("email") or None,
        "phone": phone or None,
        "street": street or None,
        "city": c.get("city") or None,
        "state": clean_state(c.get("state", "")) or None,
        "zip": clean_zip(c.get("zip", "")) or None,
        "time_preference": pref,
        "secondary_contacts": contacts,
        "notes": c.get("notes") or None,
        "area": c.get("area") or None,
        "zone": c.get("zone") or None,
    }
    if parts and parts.get("company") and (parts.get("first_name") or parts.get("last_name")):
        person = " ".join(p for p in (parts["first_name"], parts["last_name"]) if p)
        notes.append(f"{person} is saved as the contact person")
    return fields, problems, notes


# ─── the whole file ─────────────────────────────────────────────────────────


def name_key(name: Optional[str]) -> str:
    return re.sub(r"\s+", " ", (name or "").strip()).lower()


def existing_name_index(clients: Iterable[dict]) -> dict[str, str]:
    """{lowercased name or former name: the client's current name}."""
    index: dict[str, str] = {}
    for cl in clients:
        current = cl.get("name") or ""
        for n in [*(cl.get("former_names") or []), current]:
            if n:
                index[name_key(n)] = current
    return index


def validate(rows: list[list[str]], existing: dict[str, str]) -> dict:
    """The dry run: one result per non-blank data row, plus a summary.

    ``existing`` is ``existing_name_index`` over the saved clients."""
    if not rows or not any(_clean(c) for c in rows[0]):
        raise ImportFileError("The first row must be the column headers. Download the template to see them.")
    mapping, ignored = map_headers(rows[0])
    keys = set(mapping.values())
    if not keys & {"first_name", "last_name", "company", "name"}:
        raise ImportFileError(
            "Couldn't find a name column. The first row needs headers like "
            "First and Last (or Business name, or Odoo's Name).")
    data_rows = [(i + 2, r) for i, r in enumerate(rows[1:]) if any(_clean(c) for c in r)]
    if len(data_rows) > MAX_ROWS:
        raise ImportFileError(f"That file has {len(data_rows)} rows; {MAX_ROWS} at most per import.")

    results: list[dict] = []
    seen: dict[str, int] = {}
    for row_no, raw in data_rows:
        cells = {key: (raw[i] if i < len(raw) else "") for i, key in mapping.items()}
        fields, problems, notes = build_row(cells)
        name = fields["name"]
        result = {"row": row_no, "name": name, "fields": fields, "notes": notes, "message": None}
        key = name_key(name)
        if problems:
            result.update(status=STATUS_ERROR, message="; ".join(problems))
        elif key in existing:
            # The import shows names as their parts, never composed
            # (user, 2026-10-12), so the message doesn't quote the name.
            match = existing[key]
            if name_key(match) != key:
                msg = "Already a client (under a newer name)"
            elif (fields["parts"] or {}).get("company") and not fields["parts"].get("site"):
                # Businesses go by the business name alone, so a second card
                # for one business needs a Location to be told apart.
                msg = "Already a client. If this is another site, add a Location."
            else:
                msg = "Already a client"
            result.update(status=STATUS_DUPLICATE, message=msg)
        elif key in seen:
            result.update(status=STATUS_DUPLICATE, message=f"Same name as row {seen[key]} in this file")
        else:
            result["status"] = STATUS_CREATE
        if key and key not in seen and not problems:
            seen[key] = row_no
        results.append(result)

    counts = {s: sum(1 for r in results if r["status"] == s)
              for s in (STATUS_CREATE, STATUS_DUPLICATE, STATUS_ERROR)}
    return {
        "rows": results,
        "counts": {**counts, "total": len(results)},
        "columns": {rows[0][i].strip(): FIELD_LABEL.get(k, _EXTRA_LABEL.get(k, k)) for i, k in mapping.items()},
        "ignored_columns": ignored,
    }


#: How the alias-only keys read in the "columns we read" list.
_EXTRA_LABEL = {
    "name": "First + Last (or Business name)",
    "mobile": "PHONE / additional contact",
    "street2": "ADDRESS (line 2)",
    "is_company": "Business or person",
    "company_type": "Business or person",
}


def public_row(r: dict) -> dict:
    """What the API returns for a row: no internal ``fields`` dict, but the
    columns the preview table shows."""
    f = r["fields"]
    parts = f.get("parts") or {}
    return {
        "row": r["row"], "status": r["status"], "message": r["message"], "notes": r["notes"],
        "name": r["name"],
        "first_name": parts.get("first_name"), "last_name": parts.get("last_name"),
        "company": parts.get("company"), "site": parts.get("site"),
        "email": f["email"], "phone": f["phone"],
        "street": f["street"], "city": f["city"], "state": f["state"], "zip": f["zip"],
        "time_preference": f["time_preference"],
        "area": f["area"], "zone": f["zone"],
        "secondary_contacts": f["secondary_contacts"],
        "notes_text": f["notes"],
    }


def template_csv(with_examples: bool = False) -> str:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow(TEMPLATE_HEADERS)
    if with_examples:
        for ex in EXAMPLE_ROWS:
            w.writerow([ex.get(k, "") for k, _ in TEMPLATE_FIELDS])
    return out.getvalue()
