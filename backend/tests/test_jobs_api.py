"""Characterisation tests for the Jobs API (`app.apis.jobs`) and its xlsx export.

Snapshot-style: every assertion pins what the code does TODAY, including the
query shapes (N+1 loops) that later refactors are expected to change. When a
refactor changes one of these on purpose, update the test in the same commit.

Endpoints are plain coroutines driven with ``asyncio.run`` against the shared
``fake_db`` harness from ``conftest.py``.
"""

import asyncio
import io
import json
from datetime import date, datetime, timezone
from decimal import Decimal

import openpyxl
import pytest
from fastapi import HTTPException

from app.apis import jobs
from app.apis.jobs import export as jobs_export
from conftest import FakeConn

T0 = datetime(2026, 9, 1, 12, 0, tzinfo=timezone.utc)


# ─── Row builders (keys follow the SELECT column lists in jobs/__init__.py) ──


def job_row(**over):
    row = {
        "id": 1, "order_no": "MO-1", "name": "Smith Foyer", "client_id": 5, "client_name": "Smith",
        "project_id": None, "season": "2026", "collection": "Foyer", "order_date": None,
        "install_date": date(2026, 11, 20), "due_date": None, "designer": None, "sidemark": None,
        "delivery_method": None, "color_palette": "Gold", "intake": '{"phone": "555"}', "notes": None,
        "built_at": None, "installed_at": None, "created_by": "user-1", "created_at": T0, "updated_at": T0,
    }
    row.update(over)
    return row


def piece_row(**over):
    row = {"id": 20, "job_id": 1, "piece_type": "Tree", "qty": Decimal("1"), "spec": '{"height_ft": 9}',
           "design_id": None, "sort_order": 1, "created_at": T0}
    row.update(over)
    return row


def need_row(**over):
    row = {"id": 10, "job_id": 1, "piece_id": None, "label": "Sugar pine cones", "spec": "6in",
           "need_qty": Decimal("30"), "unit": "each", "shelf_qty": Decimal("4"), "source": "manual",
           "notes": None, "sort_order": 1, "created_at": T0, "updated_at": T0}
    row.update(over)
    return row


def line_row(**over):
    # l.* plus the joined PO / allocation columns from _load_job's line query.
    row = {
        "id": 100, "need_id": 10, "product_id": 555, "supplier_id": 7, "vendor_name": "Vickerman",
        "sku": "PC-20", "description": "Sugar pine cone 20pk", "image_url": None, "status": "proposed",
        "price_per": "pack", "pack_qty": 20, "covers_qty": Decimal("26"), "packs": 2,
        "order_qty": Decimal("40"), "unit_cost": Decimal("66"), "adj_unit_cost": Decimal("3.3"),
        "allocated_from_order_item_id": None, "allocated_qty": Decimal("0"), "order_item_id": None,
        "substitute_for": None, "notes": None, "created_by": "user-1", "created_at": T0, "updated_at": T0,
        "po_quantity": None, "received_qty": None, "order_id": None, "order_name": None,
        "order_status": None, "expected_arrival": None, "vendor_order_no": None, "placed_at": None,
        "allocated_order_name": None, "allocated_order_status": None, "allocated_received_qty": None,
        "allocated_line_qty": None,
    }
    row.update(over)
    return row


def task_row(**over):
    row = {"id": 3, "job_id": 1, "sourcing_line_id": None, "title": "Email Jason", "assignee": "Ana",
           "due": None, "done_at": None, "created_by": "user-1", "created_at": T0}
    row.update(over)
    return row


def seed_job(db, *, job=None, pieces=None, needs=None, lines=None, tasks=None, pos=None):
    """Register every query `_load_job` issues."""
    db.on_fetchrow("SELECT * FROM ll_app.jobs WHERE id = $1", job if job is not None else job_row())
    db.on_fetch("SELECT * FROM ll_app.job_pieces WHERE job_id", pieces if pieces is not None else [piece_row()])
    db.on_fetch("SELECT * FROM ll_app.material_needs WHERE job_id",
                needs if needs is not None else [need_row(), need_row(id=11, label="Ribbon",
                                                                      need_qty=Decimal("10"),
                                                                      shelf_qty=Decimal("10"), spec=None,
                                                                      sort_order=2)])
    db.on_fetch("FROM ll_app.sourcing_lines l JOIN ll_app.material_needs n ON n.id = l.need_id LEFT JOIN",
                lines if lines is not None else [line_row()])
    db.on_fetch("SELECT * FROM ll_app.job_tasks WHERE job_id", tasks if tasks is not None else [task_row()])
    db.on_fetch("JOIN ll_app.order_items oi ON oi.order_id = o.id AND oi.job_id = $1", pos or [])


def run(coro):
    return asyncio.run(coro)


# ─── Schema ──────────────────────────────────────────────────────────────────


def test_ensure_schema_runs_orders_and_jobs_ddl_once(fake_db):
    conn = FakeConn(fake_db)
    run(jobs.ensure_schema(conn))
    # orders DDL (one execute) + jobs DDL (one execute)
    assert fake_db.ddl_runs == 2
    assert fake_db.seen("CREATE TABLE IF NOT EXISTS ll_app.orders")
    assert fake_db.seen("CREATE TABLE IF NOT EXISTS ll_app.job_items")
    run(jobs.ensure_schema(conn))
    assert fake_db.ddl_runs == 2
    assert len(fake_db.executed) == 2


# ─── Meta / list / board-list ────────────────────────────────────────────────


def test_meta_constants():
    assert run(jobs.jobs_meta()) == {
        "stages": ["new", "sourcing", "ordered", "receiving", "complete"],
        "piece_types": [
            "Tree", "Tree Skirt", "Tree Decor", "Lighting", "Garland", "Wreath",
            "Vertical Spray", "Horizontal Swag", "Enhancers", "Low Arrangement",
            "Tall Arrangement", "Table Arrangement", "Tablescape", "Container", "Other",
        ],
        "sourcing_statuses": ["proposed", "sold_out", "ready", "ordered", "follow_up", "allocated", "on_hold"],
    }


def test_list_loads_every_job_one_by_one(fake_db):
    fake_db.on_fetch("SELECT id FROM ll_app.jobs ORDER BY updated_at DESC", [{"id": 1}, {"id": 2}])
    seed_job(fake_db)
    out = run(jobs.list_jobs())
    assert len(out) == 2
    assert list(out[0].keys()) == [
        "id", "name", "order_no", "client_name", "client_id", "project_id", "season",
        "collection", "install_date", "order_date", "due_date", "stage", "summary",
        "updated_at", "created_at",
    ]
    assert out[0]["stage"] == "sourcing"
    # N+1: one _load_job (5 fetch + 1 fetchrow) per listed job.
    assert [a for _, a in fake_db.calls("SELECT * FROM ll_app.jobs WHERE id = $1")] == [(1,), (2,)]
    assert len(fake_db.calls("FROM ll_app.sourcing_lines l")) == 2
    assert len(fake_db.executed) == 2 + 1 + 2 * 6  # DDL + id list + per-job load


def test_list_skips_ids_whose_job_row_vanished(fake_db):
    fake_db.on_fetch("SELECT id FROM ll_app.jobs ORDER BY updated_at DESC", [{"id": 9}])
    assert run(jobs.list_jobs()) == []


def test_board_list_returns_rows_verbatim(fake_db):
    rows = [{"id": 1, "name": "Smith Foyer", "client_name": "Smith", "collection": "Foyer", "season": "2026",
             "updated_at": T0, "created_at": T0, "item_count": 4, "group_count": 2, "chosen_count": 1}]
    fake_db.on_fetch("FROM ll_app.jobs j ORDER BY j.updated_at DESC", rows)
    assert run(jobs.board_list()) == rows


# ─── Create / get / patch / delete ───────────────────────────────────────────


def test_create_derives_name_and_stamps_request_user(fake_db, fake_request):
    fake_db.on_fetchrow("INSERT INTO ll_app.jobs", {"id": 1})
    seed_job(fake_db)
    body = jobs.JobCreate(client_name="Smith", collection="Foyer", intake=None)
    out = run(jobs.create_job(body, fake_request("user-7")))
    (_, args), = fake_db.calls("INSERT INTO ll_app.jobs")
    assert args == (None, "Smith Foyer", None, "Smith", None, None, "Foyer", None, None, None, None,
                    None, None, None, "{}", None, "user-7")
    assert out["id"] == 1


def test_create_untitled_and_anonymous(fake_db, fake_request):
    fake_db.on_fetchrow("INSERT INTO ll_app.jobs", {"id": 1})
    seed_job(fake_db)
    run(jobs.create_job(jobs.JobCreate(name="   "), fake_request()))
    (_, args), = fake_db.calls("INSERT INTO ll_app.jobs")
    assert args[1] == "Untitled job"
    assert args[-1] == "local-dev-user"


def test_get_404_when_missing(fake_db):
    with pytest.raises(HTTPException) as exc:
        run(jobs.get_job(99))
    assert exc.value.status_code == 404
    assert exc.value.detail == "Job not found"


def test_get_derives_need_numbers_stage_and_summary(fake_db):
    seed_job(fake_db)
    job = run(jobs.get_job(1))
    assert job["intake"] == {"phone": "555"}
    assert job["pieces"][0]["spec"] == {"height_ft": 9}
    assert job["stage"] == "sourcing"
    cones, ribbon = job["needs"]
    line = cones["lines"][0]
    assert line["unit_cost"] == 66.0 and line["covers_qty"] == 26.0 and line["order_qty"] == 40.0
    assert line["line_cost"] == 132.0  # pack price x packs
    assert line["overage_qty"] == 14.0
    assert {k: cones[k] for k in ("need_qty", "shelf_qty", "allocated_qty", "ordered_qty", "received_qty",
                                  "proposed_qty", "gap_qty", "unsourced_qty", "on_shelf_qty", "ready")} == {
        "need_qty": 30.0, "shelf_qty": 4.0, "allocated_qty": 0, "ordered_qty": 0, "received_qty": 0,
        "proposed_qty": 26.0, "gap_qty": 26.0, "unsourced_qty": 0.0, "on_shelf_qty": 4.0, "ready": False,
    }
    assert ribbon["ready"] is True and ribbon["lines"] == []
    assert job["summary"] == {"piece_count": 1, "need_count": 2, "ready_count": 1, "gap_count": 1,
                              "unsourced_count": 0, "open_tasks": 1, "buy_cost": 132.0}
    assert job["purchase_orders"] == []


@pytest.mark.parametrize(
    "needs,lines,stage",
    [
        ([], [], "new"),
        ([need_row(shelf_qty=Decimal("30"))], [], "complete"),
        ([need_row()], [line_row(order_item_id=5, status="ordered", covers_qty=Decimal("26"),
                                 received_qty=Decimal("0"))], "ordered"),
        ([need_row()], [line_row(order_item_id=5, status="ordered", covers_qty=Decimal("26"),
                                 received_qty=Decimal("10"))], "receiving"),
    ],
)
def test_get_stage_is_derived(fake_db, needs, lines, stage):
    seed_job(fake_db, needs=needs, lines=lines)
    assert run(jobs.get_job(1))["stage"] == stage


def test_patch_builds_dynamic_update(fake_db):
    seed_job(fake_db)
    body = jobs.JobUpdate(name="New", intake={"a": 1}, built=True, installed=False)
    run(jobs.update_job(1, body))
    (sql, args), = fake_db.calls("UPDATE ll_app.jobs SET")
    assert sql == ("UPDATE ll_app.jobs SET name = $1, intake = $2::jsonb, built_at = now(), "
                   "installed_at = NULL, updated_at = now() WHERE id = $3")
    assert args == ("New", '{"a": 1}', 1)


def test_patch_with_no_fields_issues_no_update(fake_db):
    seed_job(fake_db)
    run(jobs.update_job(1, jobs.JobUpdate()))
    assert not fake_db.seen("UPDATE ll_app.jobs")


def test_delete_never_404s(fake_db):
    fake_db.on_execute("DELETE FROM ll_app.jobs", "DELETE 0")
    assert run(jobs.delete_job(42)) == {"ok": True}
    assert [a for _, a in fake_db.calls("DELETE FROM ll_app.jobs")] == [(42,)]


# ─── Pieces / needs ──────────────────────────────────────────────────────────


def test_add_piece_appends_after_max_sort_order(fake_db):
    seed_job(fake_db)
    fake_db.on_fetchval("SELECT COALESCE(MAX(sort_order), 0) + 1 FROM ll_app.job_pieces", 4)
    run(jobs.add_piece(1, jobs.PieceIn(piece_type="Wreath", qty=2, spec={"diameter_in": 30})))
    (_, args), = fake_db.calls("INSERT INTO ll_app.job_pieces")
    assert args == (1, "Wreath", 2.0, '{"diameter_in": 30}', None, 4)
    assert fake_db.seen("UPDATE ll_app.jobs SET updated_at = now() WHERE id = $1")


def test_update_piece_404(fake_db):
    with pytest.raises(HTTPException) as exc:
        run(jobs.update_piece(5, jobs.PieceUpdate(qty=3)))
    assert exc.value.detail == "Piece not found"


def test_add_needs_skips_blank_labels_and_numbers_sort_order(fake_db):
    seed_job(fake_db)
    fake_db.on_fetchval("SELECT COALESCE(MAX(sort_order), 0) FROM ll_app.material_needs", 3)
    body = jobs.NeedsBulk(needs=[
        jobs.NeedIn(label=" Cones ", need_qty=30, unit=None),
        jobs.NeedIn(label="   "),
        jobs.NeedIn(label="Ribbon", need_qty=5, sort_order=99, source=None),
    ])
    run(jobs.add_needs(1, body))
    assert [a for _, a in fake_db.calls("INSERT INTO ll_app.material_needs")] == [
        (1, None, "Cones", None, 30.0, "each", 0.0, "manual", None, 4),
        (1, None, "Ribbon", None, 5.0, "each", 0.0, "manual", None, 99),
    ]


# ─── Sourcing / allocation ───────────────────────────────────────────────────


def test_add_sourcing_uses_product_snapshot_and_pack_math(fake_db, fake_request):
    seed_job(fake_db, lines=[])
    fake_db.on_fetchval("SELECT job_id FROM ll_app.material_needs WHERE id = $1", 1)
    fake_db.on_fetchrow("FROM products p LEFT JOIN suppliers s ON s.id = p.supplier_id WHERE p.id = $1", {
        "id": 555, "name": "Sugar pine cone", "supplier_sku": "PC-20", "current_price": Decimal("66.00"),
        "supplier_id": 7, "supplier_name": "Vickerman", "case_qty": None, "moq": None,
        "availability": "in_stock", "image_urls": [], "photo_url": "http://img/p.jpg",
        "raw_data": '{"BoxQty": "20"}',
    })
    fake_db.on_fetchrow("INSERT INTO ll_app.sourcing_lines", {"id": 101})
    out = run(jobs.add_sourcing(10, jobs.SourcingIn(product_id=555, price_per="pack", substitute_for=100),
                                fake_request()))
    (_, args), = fake_db.calls("INSERT INTO ll_app.sourcing_lines")
    assert args == (10, 555, 7, "Vickerman", "PC-20", "Sugar pine cone", "http://img/p.jpg", "proposed",
                    "pack", 20, 26.0, 2, 40, 66.0, 3.3, 100, None, "local-dev-user")
    assert [a for _, a in fake_db.calls("SET status = 'sold_out'")] == [(100,)]
    assert out["created_line_id"] == 101


def test_add_sourcing_rejects_unknown_status(fake_db, fake_request):
    seed_job(fake_db)
    fake_db.on_fetchval("SELECT job_id FROM ll_app.material_needs WHERE id = $1", 1)
    with pytest.raises(HTTPException) as exc:
        run(jobs.add_sourcing(10, jobs.SourcingIn(vendor_name="X", status="bogus"), fake_request()))
    assert exc.value.status_code == 400
    assert exc.value.detail == "Unknown sourcing status"


def _alloc_src(**over):
    row = {"id": 500, "order_id": 60, "product_id": 555, "quantity": 48, "variant_note": None,
           "unit_price": None, "name_snapshot": None, "sku_snapshot": "SNAP-1", "supplier_snapshot": None,
           "added_by": None, "created_at": T0, "job_id": None, "need_id": None, "sourcing_line_id": None,
           "received_qty": Decimal("0"), "follow_up_note": None, "order_name": "Market order",
           "pname": "Hydrangea", "supplier_sku": "HY-1", "supplier_id": 7, "supplier_name": "Vickerman",
           "current_price": Decimal("4.25"), "image_urls": None, "photo_url": "http://img/h.jpg"}
    row.update(over)
    return row


def test_allocate_refuses_more_than_remaining(fake_db, fake_request):
    fake_db.on_fetchval("SELECT job_id FROM ll_app.material_needs WHERE id = $1", 1)
    fake_db.on_fetchrow("FROM ll_app.order_items oi JOIN ll_app.orders o ON o.id = oi.order_id", _alloc_src())
    fake_db.on_fetchval("SELECT COALESCE(SUM(allocated_qty), 0) FROM ll_app.sourcing_lines", Decimal("5"))
    with pytest.raises(HTTPException) as exc:
        run(jobs.allocate(10, jobs.AllocateIn(order_item_id=500, qty=44), fake_request()))
    assert exc.value.status_code == 400
    assert exc.value.detail == "Only 43 left unallocated on that order line"


def test_allocate_inserts_allocated_line(fake_db, fake_request):
    seed_job(fake_db)
    fake_db.on_fetchval("SELECT job_id FROM ll_app.material_needs WHERE id = $1", 1)
    fake_db.on_fetchrow("FROM ll_app.order_items oi JOIN ll_app.orders o ON o.id = oi.order_id", _alloc_src())
    fake_db.on_fetchval("SELECT COALESCE(SUM(allocated_qty), 0) FROM ll_app.sourcing_lines", Decimal("5"))
    run(jobs.allocate(10, jobs.AllocateIn(order_item_id=500, qty=12), fake_request("user-7")))
    (sql, args), = fake_db.calls("INSERT INTO ll_app.sourcing_lines")
    assert "VALUES ($1,$2,$3,$4,$5,$6,$7,'allocated','each',1,$8,0,0,$9,$9,$10,$8,$11,$12)" in sql
    assert args == (10, 555, 7, "Vickerman", "SNAP-1", "Hydrangea", "http://img/h.jpg", 12.0, 4.25, 500,
                    "Allocated from Market order", "user-7")


def test_allocate_line_bought_for_a_job_is_fully_reserved(fake_db, fake_request):
    fake_db.on_fetchval("SELECT job_id FROM ll_app.material_needs WHERE id = $1", 1)
    fake_db.on_fetchrow("FROM ll_app.order_items oi JOIN ll_app.orders o ON o.id = oi.order_id",
                        _alloc_src(job_id=3))
    with pytest.raises(HTTPException) as exc:
        run(jobs.allocate(10, jobs.AllocateIn(order_item_id=500, qty=1), fake_request()))
    assert exc.value.detail == "Only 0 left unallocated on that order line"


# ─── Send to PO ──────────────────────────────────────────────────────────────


def test_send_to_po_issues_per_line_statements(fake_db, fake_request):
    """Pins the N+1 shape: 1 PO insert per vendor, then 4 statements per line."""
    lines = [
        line_row(id=100, product_id=555, status="ready"),
        line_row(id=101, product_id=556, status="follow_up", price_per="each", unit_cost=Decimal("2.5"),
                 order_qty=Decimal("12.5"), covers_qty=Decimal("12")),
    ]
    seed_job(fake_db, lines=lines)
    fake_db.on_fetchval("INSERT INTO ll_app.orders", 900)
    item_ids = iter([7001, 7002])
    fake_db.on("INSERT INTO ll_app.order_items", lambda sql, *a: next(item_ids), method="fetchval")

    out = run(jobs.send_to_po(1, jobs.SendToPO(), fake_request("user-7")))

    assert out["created_orders"] == [900]
    assert [a for _, a in fake_db.calls("INSERT INTO ll_app.orders")] == [
        ("Smith Foyer · Vickerman", "Created from job #1", "user-7", 7)]
    assert [a for _, a in fake_db.calls("SELECT id FROM ll_app.order_items WHERE order_id")] == [
        (900, 555), (900, 556)]
    assert [a for _, a in fake_db.calls("INSERT INTO ll_app.order_items")] == [
        (900, 555, 40, None, 3.3, "Sugar pine cone 20pk", "PC-20", "Vickerman", "user-7", 1, 10, 100),
        (900, 556, 13, None, 2.5, "Sugar pine cone 20pk", "PC-20", "Vickerman", "user-7", 1, 10, 101),
    ]
    assert [a for _, a in fake_db.calls("UPDATE ll_app.sourcing_lines SET order_item_id")] == [
        (100, 7001), (101, 7002)]
    assert len(fake_db.calls("UPDATE ll_app.orders SET updated_at = now()")) == 2
    # Two _load_job passes (before + after) bracket the loop.
    assert len(fake_db.calls("SELECT * FROM ll_app.jobs WHERE id = $1")) == 2
    inserts = [s for s, _ in fake_db.executed if s.lstrip().upper().startswith("INSERT")]
    selects = [s for s, _ in fake_db.executed if s.lstrip().upper().startswith("SELECT")]
    assert len(inserts) == 3
    assert len(selects) == 2 * 6 + 2


def test_send_to_po_400_when_nothing_ready(fake_db, fake_request):
    seed_job(fake_db, lines=[line_row(status="sold_out")])
    with pytest.raises(HTTPException) as exc:
        run(jobs.send_to_po(1, jobs.SendToPO(), fake_request()))
    assert exc.value.detail == "No sourcing lines are ready to send"


# ─── Tasks ───────────────────────────────────────────────────────────────────


def test_add_task_strips_title(fake_db, fake_request):
    seed_job(fake_db)
    run(jobs.add_task(1, jobs.TaskIn(title="  Call Jason  ", assignee="Ana"), fake_request()))
    (_, args), = fake_db.calls("INSERT INTO ll_app.job_tasks")
    assert args == (1, None, "Call Jason", "Ana", None, "local-dev-user")


def test_update_task_done_sets_done_at(fake_db):
    seed_job(fake_db)
    fake_db.on_fetchval("SELECT job_id FROM ll_app.job_tasks WHERE id = $1", 1)
    run(jobs.update_task(3, jobs.TaskUpdate(title="X", done=True)))
    (sql, args), = fake_db.calls("UPDATE ll_app.job_tasks SET")
    assert sql == "UPDATE ll_app.job_tasks SET title = $1, done_at = now() WHERE id = $2"
    assert args == ("X", 3)


def test_delete_task_404(fake_db):
    with pytest.raises(HTTPException) as exc:
        run(jobs.delete_task(3))
    assert exc.value.detail == "Task not found"


# ─── Groups / pinned items / board ───────────────────────────────────────────


def board_item_row(**over):
    row = {
        "item_id": 1, "job_id": 1, "group_id": 2, "product_id": 555, "note": None,
        "qty_needed": Decimal("12"), "chosen": True, "sort_order": 1, "added_by": "user-1", "pinned_at": T0,
        "name": "Cone", "supplier_sku": "PC", "current_price": Decimal("3.5"), "image_urls": ["a", "b"],
        "photo_url": "a", "height_in": Decimal("6"), "width_in": None, "length_in": None,
        "diameter_in": None, "color": "brown", "finish": None, "material": None, "style": None,
        "category": "Pinecones", "case_qty": None, "moq": None, "uom": None, "unit": None,
        "availability": None,
        "raw_data": json.dumps({"BoxQty": " 20 ", "normalized": {"color": "Brown", "size_in": 6},
                                "source_photo_url": "c", "product_type": "Cone", "url": "http://x"}),
        "supplier_id": 7, "supplier_name": "Vickerman",
    }
    row.update(over)
    return row


def seed_board(db, items=None):
    db.on_fetchrow("SELECT id, name, client_name, client_id, collection, season, notes, updated_at FROM ll_app.jobs",
                   {"id": 1, "name": "Smith Foyer", "client_name": "Smith", "client_id": 5,
                    "collection": "Foyer", "season": "2026", "notes": None, "updated_at": T0})
    db.on_fetch("SELECT id, name, sort_order FROM ll_app.job_groups", [{"id": 2, "name": "Ornaments", "sort_order": 1}])
    db.on_fetch("FROM ll_app.job_items i LEFT JOIN products p", items if items is not None else [board_item_row()])


def test_board_normalises_items(fake_db):
    seed_board(fake_db, items=[board_item_row(), board_item_row(item_id=2, product_id=9, name=None,
                                                                 image_urls=None, photo_url=None,
                                                                 raw_data=None, qty_needed=None,
                                                                 current_price=None, height_in=None)])
    board = run(jobs.get_board(1))
    assert board["groups"] == [{"id": 2, "name": "Ornaments", "sort_order": 1}]
    first, missing = board["items"]
    assert "raw_data" not in first
    assert {k: first[k] for k in ("image_urls", "box_qty", "current_price", "height_in", "qty_needed",
                                  "product_url", "norm_color", "norm_finish", "norm_size_in",
                                  "product_type", "missing")} == {
        "image_urls": ["a", "b", "c"], "box_qty": 20, "current_price": 3.5, "height_in": 6.0,
        "qty_needed": 12.0, "product_url": "http://x", "norm_color": "Brown", "norm_finish": None,
        "norm_size_in": 6, "product_type": "Cone", "missing": False,
    }
    assert missing["missing"] is True and missing["image_urls"] == [] and missing["box_qty"] is None


def test_board_404(fake_db):
    with pytest.raises(HTTPException) as exc:
        run(jobs.get_board(1))
    assert exc.value.detail == "Job not found"


def test_add_group_requires_name_and_returns_created_id(fake_db):
    with pytest.raises(HTTPException) as exc:
        run(jobs.add_group(1, jobs.GroupIn(name="  ")))
    assert exc.value.status_code == 400
    seed_board(fake_db)
    fake_db.on_fetchval("SELECT COALESCE(MAX(sort_order), 0) + 1 FROM ll_app.job_groups", 3)
    fake_db.on_fetchval("INSERT INTO ll_app.job_groups", 44)
    out = run(jobs.add_group(1, jobs.GroupIn(name=" Garland ")))
    assert [a for _, a in fake_db.calls("INSERT INTO ll_app.job_groups")] == [(1, "Garland", 3)]
    assert out["created_group_id"] == 44


def test_pin_item_creates_then_moves_existing(fake_db, fake_request):
    fake_db.on_fetchval("SELECT 1 FROM ll_app.jobs WHERE id = $1", 1)
    fake_db.on_fetchval("SELECT 1 FROM products WHERE id = $1", 1)
    fake_db.on_fetchval("SELECT COALESCE(MAX(sort_order), 0) + 1 FROM ll_app.job_items", 2)
    fake_db.on_fetchval("INSERT INTO ll_app.job_items", 77)
    fake_db.on_fetch("SELECT id, name FROM ll_app.job_groups", [])
    fake_db.on_fetch("SELECT product_id, group_id FROM ll_app.job_items", [{"product_id": 555, "group_id": None}])
    out = run(jobs.pin_item(1, jobs.PinIn(product_id=555), fake_request()))
    assert out == {"item_id": 77, "created": True, "job_id": 1, "groups": [],
                   "pins": [{"product_id": 555, "group_id": None}]}
    assert [a for _, a in fake_db.calls("INSERT INTO ll_app.job_items")] == [(1, None, 555, None, 2, "local-dev-user")]

    fake_db.on_fetchrow("SELECT id FROM ll_app.job_items WHERE job_id = $1 AND product_id = $2", {"id": 77})
    fake_db.on_fetchval("SELECT 1 FROM ll_app.job_groups WHERE id = $1 AND job_id = $2", 1)
    out = run(jobs.pin_item(1, jobs.PinIn(product_id=555, group_id=2, note="hero"), fake_request()))
    assert out["created"] is False and out["item_id"] == 77
    assert [a for _, a in fake_db.calls("UPDATE ll_app.job_items SET group_id = COALESCE")] == [(77, 2, "hero")]


def test_pin_item_rejects_group_from_another_job(fake_db, fake_request):
    fake_db.on_fetchval("SELECT 1 FROM ll_app.jobs WHERE id = $1", 1)
    fake_db.on_fetchval("SELECT 1 FROM products WHERE id = $1", 1)
    with pytest.raises(HTTPException) as exc:
        run(jobs.pin_item(1, jobs.PinIn(product_id=555, group_id=99), fake_request()))
    assert exc.value.status_code == 400
    assert exc.value.detail == "That group is not on this job"


def test_update_pin_clear_group_and_fields(fake_db):
    seed_board(fake_db)
    fake_db.on_fetchval("SELECT job_id FROM ll_app.job_items WHERE id = $1", 1)
    run(jobs.update_pin(1, jobs.PinUpdate(clear_group=True, group_id=5, qty_needed=3, chosen=False)))
    (sql, args), = fake_db.calls("UPDATE ll_app.job_items SET")
    assert sql == "UPDATE ll_app.job_items SET group_id = NULL, qty_needed = $1, chosen = $2 WHERE id = $3"
    assert args == (3.0, False, 1)


# ─── Export ──────────────────────────────────────────────────────────────────


def test_export_xlsx_sheets_and_rows(fake_db):
    seed_job(fake_db)
    resp = run(jobs.export_job(1))
    assert resp.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert resp.headers["content-disposition"] == 'attachment; filename="Smith_Foyer_tracking.xlsx"'
    wb = openpyxl.load_workbook(io.BytesIO(resp.body))
    assert wb.sheetnames == ["Order", "Request"]
    rows = list(wb["Order"].iter_rows(values_only=True))
    assert rows[0] == ("Client", "Project", "Item", "Vendor", "Sku", "Description", "Need Qty", "O/O QTY",
                       "Unit Cost", "Adj. Unit Cost", "Freight", "Arrival", "Checked in", "Notes", "Picture")
    assert rows[1][0] == "Not yet ordered"
    assert rows[2] == ("Smith", "Foyer", "Sugar pine cones", "Vickerman", "PC-20", "Sugar pine cone 20pk",
                       30, 40, 66, 3.3, None, None, None, "14 extra to stock", None)
    assert rows[3] == ("Smith", "Foyer", "Ribbon", None, None, None, 10, None, None, None, None, None, None,
                       None, None)
    assert list(wb["Request"].iter_rows(values_only=True)) == [
        (None, "Color Scheme : Foyer"),
        (1, "Tree : 1  (height_ft: 9)"),
        (None, "Phone : 555"),
    ]


def test_export_rejects_non_xlsx_after_loading(fake_db):
    seed_job(fake_db)
    with pytest.raises(HTTPException) as exc:
        run(jobs.export_job(1, format="pdf"))
    assert exc.value.status_code == 400
    assert exc.value.detail == "format must be xlsx"
    assert fake_db.seen("SELECT * FROM ll_app.jobs WHERE id = $1")  # loaded before the format check


def test_export_404(fake_db):
    with pytest.raises(HTTPException) as exc:
        run(jobs.export_job(1))
    assert exc.value.status_code == 404


# ─── export.py helpers ───────────────────────────────────────────────────────


@pytest.mark.parametrize(
    "value,expected",
    [(0, "$0.00"), (1234.5, "$1,234.50"), (1234.567, "$1,234.57"), (None, ""), (-5, "$-5.00")],
)
def test_money(value, expected):
    assert jobs_export._money(value) == expected


@pytest.mark.parametrize("value", ["abc", '<a href="x">&</a>'])
def test_money_raises_on_non_numeric_strings(value):
    with pytest.raises(ValueError):
        jobs_export._money(value)


@pytest.mark.parametrize(
    "value,expected",
    [
        (0, ""),  # falsy zero renders as empty
        (1234.5, "1234.5"),
        (1234.567, "1234.567"),
        (None, ""),
        ("abc", "abc"),
        (-5, "-5"),
        ('<a href="x">&</a>', '&lt;a href="x"&gt;&amp;&lt;/a&gt;'),  # double quotes NOT escaped
    ],
)
def test_esc(value, expected):
    assert jobs_export._esc(value) == expected
